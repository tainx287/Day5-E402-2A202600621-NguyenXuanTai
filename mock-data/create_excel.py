import json
import os
import sys

# Configure stdout encoding to UTF-8
if sys.version_info >= (3, 7):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Thư viện 'openpyxl' chưa được cài đặt. Vui lòng chạy: pip install openpyxl")
    sys.exit(1)

JSON_FILE = os.path.join(os.path.dirname(__file__), "mock_hotel_database.json")
EXCEL_FILE = os.path.join(os.path.dirname(__file__), "mock_hotel_database.xlsx")

def load_data():
    with open(JSON_FILE, "r", encoding="utf-8") as f:
        return json.load(f)

def build_excel():
    db = load_data()
    wb = openpyxl.Workbook()
    
    # Định nghĩa style chung
    font_family = "Segoe UI"
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    
    title_font = Font(name=font_family, size=14, bold=True, color="1F4E78")
    
    cell_font = Font(name=font_family, size=10)
    bold_cell_font = Font(name=font_family, size=10, bold=True)
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='center')

    # 1. SHEET 1: Hotel Info
    ws_info = wb.active
    ws_info.title = "Hotel General Info"
    ws_info.views.sheetView[0].showGridLines = True
    
    ws_info.append([])
    ws_info.cell(row=2, column=2, value="THÔNG TIN CHUNG KHÁCH SẠN").font = title_font
    ws_info.append([])
    
    info_data = db.get("hotel_info", {})
    general_fields = [
        ("Tên khách sạn", info_data.get("hotel_name")),
        ("Mô tả ngắn", info_data.get("short_description")),
        ("Mô tả đầy đủ", info_data.get("full_description")),
        ("Địa chỉ", info_data.get("address")),
        ("Khu vực", info_data.get("area")),
        ("Giờ nhận phòng (Checkin)", info_data.get("checkin_time")),
        ("Giờ trả phòng (Checkout)", info_data.get("checkout_time")),
        ("Chính sách chung", info_data.get("hotel_policies_known")),
        ("Đường dẫn ảnh chung", "; ".join(info_data.get("general_images", [])))
    ]
    
    ws_info.cell(row=4, column=2, value="Thuộc tính").font = header_font
    ws_info.cell(row=4, column=2).fill = header_fill
    ws_info.cell(row=4, column=2).alignment = align_center
    ws_info.cell(row=4, column=2).border = thin_border
    
    ws_info.cell(row=4, column=3, value="Giá trị dữ liệu").font = header_font
    ws_info.cell(row=4, column=3).fill = header_fill
    ws_info.cell(row=4, column=3).alignment = align_center
    ws_info.cell(row=4, column=3).border = thin_border
    
    current_row = 5
    for field, val in general_fields:
        c1 = ws_info.cell(row=current_row, column=2, value=field)
        c2 = ws_info.cell(row=current_row, column=3, value=val)
        
        c1.font = bold_cell_font
        c1.border = thin_border
        c1.alignment = Alignment(horizontal='left', vertical='top')
        
        c2.font = cell_font
        c2.border = thin_border
        c2.alignment = align_left
        
        current_row += 1
        
    ws_info.column_dimensions['B'].width = 25
    ws_info.column_dimensions['C'].width = 85

    # 2. SHEET 2: Hotel Amenities
    ws_amenities = wb.create_sheet(title="Hotel Amenities")
    ws_amenities.views.sheetView[0].showGridLines = True
    
    headers_am = ["Tên tiện ích", "Khả dụng", "Miễn phí", "Ghi chú chi phí", "Áp dụng cho phòng", "Nguồn dữ liệu"]
    ws_amenities.append(headers_am)
    
    for col_idx, h in enumerate(headers_am, 1):
        cell = ws_amenities.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border
        
    for amenity in info_data.get("general_amenities", []):
        row_data = [
            amenity.get("amenity_name"),
            "Có" if amenity.get("available") else "Không",
            "Miễn phí" if amenity.get("is_free") else "Tính phí",
            amenity.get("fee_note"),
            amenity.get("applies_to_room"),
            amenity.get("source")
        ]
        ws_amenities.append(row_data)
        
    for row in range(2, len(info_data.get("general_amenities", [])) + 2):
        for col in range(1, 7):
            cell = ws_amenities.cell(row=row, column=col)
            cell.font = cell_font
            cell.border = thin_border
            if col in [2, 3, 6]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    # 3. SHEET 3: Room Types
    ws_rooms = wb.create_sheet(title="Room Types")
    ws_rooms.views.sheetView[0].showGridLines = True
    
    headers_rooms = [
        "Mã phòng", "Tên hạng phòng", "Mô tả phòng", "Giá từ (VNĐ)", 
        "Sức chứa (Khách)", "Diện tích (m²)", "Loại giường", 
        "Tầm nhìn (View)", "Tiện ích trong phòng", "Ảnh phòng", "Ghi chú thêm"
    ]
    ws_rooms.append(headers_rooms)
    
    for col_idx, h in enumerate(headers_rooms, 1):
        cell = ws_rooms.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border
        
    rooms_data = db.get("room_types", [])
    for r in rooms_data:
        row_data = [
            r.get("room_id"),
            r.get("room_name"),
            r.get("room_description"),
            r.get("price_from"),
            r.get("max_guests"),
            r.get("area_m2"),
            r.get("bed_type"),
            r.get("view"),
            "\n".join([f"• {a}" for a in r.get("room_amenities", [])]),
            "; ".join([f"{img.get('caption', '')} ({img.get('url', '')})" for img in r.get("room_images", [])]),
            r.get("notes")
        ]
        ws_rooms.append(row_data)
        
    for r_idx in range(2, len(rooms_data) + 2):
        for col in range(1, 12):
            cell = ws_rooms.cell(row=r_idx, column=col)
            cell.font = cell_font
            cell.border = thin_border
            if col in [1, 5, 6]:
                cell.alignment = align_center
            elif col == 4:
                cell.alignment = align_right
                cell.number_format = '#,##0'
            else:
                cell.alignment = align_left

    # 4. SHEET 4: Nearby Places
    ws_places = wb.create_sheet(title="Nearby Places")
    ws_places.views.sheetView[0].showGridLines = True
    
    headers_places = ["Mã địa điểm", "Tên địa điểm", "Phân loại", "Khoảng cách", "Mô tả chi tiết", "Phù hợp cho", "Nguồn xác minh"]
    ws_places.append(headers_places)
    
    for col_idx, h in enumerate(headers_places, 1):
        cell = ws_places.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border
        
    places_data = db.get("nearby_places", [])
    for p in places_data:
        cat_map = {
            "attraction": "Điểm tham quan",
            "restaurant": "Nhà hàng/Café",
            "atm": "ATM",
            "convenience_store": "Cửa hàng tiện lợi",
            "pharmacy": "Nhà thuốc"
        }
        cat_vi = cat_map.get(p.get("category"), p.get("category"))
        row_data = [
            p.get("place_id"),
            p.get("place_name"),
            cat_vi,
            p.get("distance"),
            p.get("short_description"),
            p.get("suitable_for"),
            p.get("source_note")
        ]
        ws_places.append(row_data)
        
    for r_idx in range(2, len(places_data) + 2):
        for col in range(1, 8):
            cell = ws_places.cell(row=r_idx, column=col)
            cell.font = cell_font
            cell.border = thin_border
            if col in [1, 3, 4, 7]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    # 5. SHEET 5: Policy & FAQ
    ws_faq = wb.create_sheet(title="Policy & FAQ")
    ws_faq.views.sheetView[0].showGridLines = True
    
    headers_faq = ["Câu hỏi thường gặp", "Câu trả lời chính thức", "Phân loại", "Nguồn tài liệu", "Độ tin cậy"]
    ws_faq.append(headers_faq)
    
    for col_idx, h in enumerate(headers_faq, 1):
        cell = ws_faq.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border
        
    faq_data = db.get("policy_or_faq", [])
    for faq in faq_data:
        row_data = [
            faq.get("question"),
            faq.get("answer"),
            faq.get("category"),
            faq.get("source"),
            faq.get("confidence")
        ]
        ws_faq.append(row_data)
        
    for r_idx in range(2, len(faq_data) + 2):
        for col in range(1, 6):
            cell = ws_faq.cell(row=r_idx, column=col)
            cell.font = cell_font
            cell.border = thin_border
            if col in [3, 4, 5]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    # Tự động điều chỉnh độ rộng cột cho các sheet có cấu trúc bảng
    for ws in [ws_amenities, ws_rooms, ws_places, ws_faq]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            
            for cell in col:
                val_str = str(cell.value or '')
                # Nếu chuỗi có xuống dòng, lấy dòng dài nhất làm chuẩn độ rộng
                lines = val_str.split('\n')
                for line in lines:
                    if len(line) > max_len:
                        max_len = len(line)
            
            # Giới hạn độ rộng tối đa để cột không quá bè
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)
            
    # Lưu file Excel
    wb.save(EXCEL_FILE)
    print(f"Xuat Excel thanh cong tai: {EXCEL_FILE}")

if __name__ == "__main__":
    build_excel()
